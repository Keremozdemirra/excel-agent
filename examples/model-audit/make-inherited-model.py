#!/usr/bin/env python3
"""Rebuild inherited-model.xlsx, the workbook used in the worked example.

The model is a twelve-month subscription plan of the shape an operator
would actually hand over: an assumptions sheet, a monthly build, a cost
block, a summary, and a scratch sheet somebody forgot to delete. Every
defect the audit reports was planted deliberately and is listed in
notes.md.
"""

from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName

MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun",
          "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]

wb = Workbook()

# ---------------------------------------------------------------- Inputs
inp = wb.active
inp.title = "Inputs"
inp["A1"] = "FY2025 plan — assumptions"
inp["A2"], inp["B2"] = "Assumption", "Value"
rows = [
    ("Opening customers", 1200),
    ("Monthly new customers", 85),
    ("Monthly churn rate", 0.021),
    ("ARPU (TRY/month)", 480),
    ("Gross margin", 0.62),
    ("Headcount", 6),
    ("Average salary (TRY/month)", 48000),
    ("Marketing (TRY/month)", 60000),
    ("Office and other (TRY/month)", 35000),
    ("FX rate TRY/EUR", 38.4),
]
for i, (label, value) in enumerate(rows, start=3):
    inp[f"A{i}"] = label
    inp[f"B{i}"] = value

# --------------------------------------------------------------- Revenue
rev = wb.create_sheet("Revenue")
rev["A1"] = "Monthly build"
rev["A3"] = "Month"
for i, m in enumerate(MONTHS):
    rev.cell(row=3, column=2 + i, value=m)

rev["A5"] = "Opening customers"
rev["A6"] = "New customers"
rev["A7"] = "Churned customers"
rev["A9"] = "Ending customers"
rev["A11"] = "ARPU"
rev["A12"] = "Revenue"
rev["A13"] = "Gross profit"
rev["A15"] = "Total revenue"

for i in range(12):
    col = get_column_letter(2 + i)
    prev = get_column_letter(1 + i)
    rev[f"{col}5"] = "=Inputs!$B$3" if i == 0 else f"={prev}9"
    rev[f"{col}6"] = "=Inputs!$B$4"
    rev[f"{col}7"] = f"=-{col}5*Inputs!$B$5"
    rev[f"{col}9"] = f"={col}5+{col}6+{col}7"
    rev[f"{col}11"] = "=Inputs!$B$6"
    rev[f"{col}12"] = f"={col}9*{col}11"
    rev[f"{col}13"] = f"={col}12*Inputs!$B$7"

# planted: price rise typed into the formula instead of an input cell
rev["F11"] = "=Inputs!$B$6*1.06"
# planted: a month overwritten with a typed number
rev["H12"] = 762000
# planted: a different margin in one month, typed inline
rev["K13"] = "=K12*0.55"
# planted: annual total stops at October
rev["B15"] = "=SUM(B12:K12)"

# ----------------------------------------------------------------- Costs
cost = wb.create_sheet("Costs")
cost["A1"] = "Operating cost"
cost["A3"], cost["B3"] = "Cost line", "TRY per month"
cost["A4"], cost["B4"] = "Salaries", "=Inputs!$B$8*Inputs!$B$9"
cost["A5"], cost["B5"] = "Marketing", "=Inputs!$B$10"
cost["A6"], cost["B6"] = "Office and other", "=Inputs!$B$11"
# planted: per-customer hosting cost buried in the formula
cost["A7"], cost["B7"] = "Cloud hosting", "=Inputs!$B$3*13"
# planted: per-seat tooling cost buried in the formula
cost["A8"], cost["B8"] = "Tooling per seat", "=Inputs!$B$8*1800"
# planted: per-seat licence price in EUR buried in the formula
cost["A9"], cost["B9"] = "Software licences", "=Inputs!$B$8*42*Inputs!$B$12"
# planted: travel as a share of marketing, buried in the formula
cost["A10"], cost["B10"] = "Travel", "=Inputs!$B$10*0.08"
cost["A11"], cost["B11"] = "Recruitment", 18000
cost["A12"], cost["B12"] = "Legal and audit", 12000
# planted: total misses the last line item
cost["A13"], cost["B13"] = "Total monthly cost", "=SUM(B4:B11)"
cost["A15"], cost["B15"] = "Annual cost", "=B13*12"

# --------------------------------------------------------------- Summary
sm = wb.create_sheet("Summary")
sm["A1"] = "FY2025 summary"
sm["A3"], sm["B3"] = "Revenue", "=Revenue!B15"
sm["A4"], sm["B4"] = "Gross profit", "=SUM(Revenue!B13:M13)"
sm["A5"], sm["B5"] = "Operating cost", "=Costs!B15"
sm["A6"], sm["B6"] = "EBITDA", "=B4-B5"
sm["A8"], sm["B8"] = "EBITDA margin", "=B6/B3"
sm["A10"], sm["B10"] = "Ending customers", "=Revenue!M9"
sm["A11"], sm["B11"] = "Revenue per head", "=B3/Inputs!$B$8"
# planted: link to last year's plan, which nobody has a copy of
sm["A13"], sm["B13"] = "Prior-year revenue", "='[Plan_2024.xlsx]Summary'!$B$3"
sm["A14"], sm["B14"] = "Growth vs prior year", "=B3/B13-1"
# planted: survivor of a deleted sheet
sm["A16"], sm["B16"] = "Legacy check", "=#REF!*1.15"

# --------------------------------------------------------------- Scratch
sc = wb.create_sheet("Scratch")
sc["A1"] = "Old upside case — do not use"
sc["A3"] = "Month"
for i in range(3):
    col = get_column_letter(2 + i)
    src = get_column_letter(2 + i)
    sc.cell(row=3, column=2 + i, value=MONTHS[i])
    for r in range(4, 10):
        sc[f"{col}{r}"] = f"=Revenue!{src}12-Revenue!{src}13"

try:
    wb.defined_names["Eski_Ciro"] = DefinedName("Eski_Ciro", attr_text="Revenue!#REF!")
except Exception:  # older openpyxl API
    pass

wb.save("inherited-model.xlsx")
print("wrote inherited-model.xlsx")
